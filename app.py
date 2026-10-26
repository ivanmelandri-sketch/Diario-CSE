import os
import io
import time
from datetime import datetime

import pytz
import requests
import psycopg2
from flask import Flask, render_template_string, request, redirect, url_for, send_file

app = Flask(__name__)

TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN")
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY")
GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.8-flash")
GEMINI_FALLBACK_MODEL = os.environ.get("GEMINI_FALLBACK_MODEL", "gemini-3.8-flash-lite")
DATABASE_URL = os.environ.get("DATABASE_URL")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ITALY_TZ = pytz.timezone("Europe/Rome")

# Note in attesa di conferma, per chat: {"text", "operator", "awaiting_edit"}
PENDING_EDITS = {}

GEMINI_RETRY_EVERY = 3     # secondi tra un tentativo e l'altro
GEMINI_MAX_WAIT = 15       # secondi massimi per ciascun modello (principale, poi riserva)


# ---------------------------------------------------------------- DATABASE

def get_db_connection():
    if DATABASE_URL:
        conn = psycopg2.connect(DATABASE_URL, sslmode="require")
        return conn, "postgres"
    import sqlite3
    conn = sqlite3.connect(os.path.join(BASE_DIR, "diario.db"))
    return conn, "sqlite"


def q(sql, db_type):
    """Adatta i segnaposto: %s per postgres, ? per sqlite."""
    return sql.replace("?", "%s") if db_type == "postgres" else sql


def init_db():
    conn, db_type = get_db_connection()
    cur = conn.cursor()
    if db_type == "postgres":
        cur.execute("CREATE TABLE IF NOT EXISTS entries (id SERIAL PRIMARY KEY, timestamp TEXT NOT NULL, operator TEXT NOT NULL, text TEXT NOT NULL)")
        cur.execute("CREATE TABLE IF NOT EXISTS operators (chat_id BIGINT PRIMARY KEY, name TEXT NOT NULL, registered_at TEXT NOT NULL)")
    else:
        cur.execute("CREATE TABLE IF NOT EXISTS entries (id INTEGER PRIMARY KEY AUTOINCREMENT, timestamp TEXT NOT NULL, operator TEXT NOT NULL, text TEXT NOT NULL)")
        cur.execute("CREATE TABLE IF NOT EXISTS operators (chat_id INTEGER PRIMARY KEY, name TEXT NOT NULL, registered_at TEXT NOT NULL)")
    conn.commit()
    conn.close()


init_db()


def now_str():
    return datetime.now(ITALY_TZ).strftime("%d/%m/%Y %H:%M")


def register_operator(chat_id, name):
    """Registra l'operatore alla prima nota; se esiste già aggiorna il nome."""
    conn, db_type = get_db_connection()
    cur = conn.cursor()
    try:
        if db_type == "postgres":
            cur.execute(
                "INSERT INTO operators (chat_id, name, registered_at) VALUES (%s, %s, %s) "
                "ON CONFLICT (chat_id) DO UPDATE SET name = EXCLUDED.name",
                (chat_id, name, now_str()),
            )
        else:
            cur.execute(
                "INSERT INTO operators (chat_id, name, registered_at) VALUES (?, ?, ?) "
                "ON CONFLICT(chat_id) DO UPDATE SET name = excluded.name",
                (chat_id, name, now_str()),
            )
        conn.commit()
    except Exception:
        app.logger.exception("Errore registrazione operatore")
    finally:
        conn.close()


def save_entry(timestamp, operator, text):
    conn, db_type = get_db_connection()
    cur = conn.cursor()
    cur.execute(q("INSERT INTO entries (timestamp, operator, text) VALUES (?, ?, ?)", db_type),
                (timestamp, operator, text))
    conn.commit()
    conn.close()


def load_entries(search_query=None, operator=None):
    conn, db_type = get_db_connection()
    cur = conn.cursor()
    like = "ILIKE" if db_type == "postgres" else "LIKE"
    sql = "SELECT id, timestamp, operator, text FROM entries WHERE 1=1"
    params = []
    if search_query:
        sql += f" AND text {like} ?"
        params.append(f"%{search_query}%")
    if operator:
        sql += " AND operator = ?"
        params.append(operator)
    sql += " ORDER BY id DESC"
    cur.execute(q(sql, db_type), params)
    rows = cur.fetchall()
    conn.close()
    return [{"id": r[0], "timestamp": r[1], "operator": r[2], "text": r[3]} for r in rows]


def load_operator_names():
    conn, _ = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT DISTINCT operator FROM entries ORDER BY operator")
    names = [r[0] for r in cur.fetchall()]
    conn.close()
    return names


def delete_entry(entry_id):
    conn, db_type = get_db_connection()
    cur = conn.cursor()
    cur.execute(q("DELETE FROM entries WHERE id = ?", db_type), (entry_id,))
    conn.commit()
    conn.close()


# ------------------------------------------------------------------ GEMINI

def call_gemini_once(text, model, timeout=5):
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    headers = {"Content-Type": "application/json", "x-goog-api-key": GEMINI_API_KEY}
    prompt = ("Sei un assistente di redazione per un team socio-educativo. "
              "Trasforma gli appunti in un paragrafo descrittivo chiaro, sintetico e professionale, "
              "senza data, ora, firme o saluti.")
    payload = {"contents": [{"parts": [{"text": f"{prompt}\n\nTesto:\n{text}"}]}]}
    r = requests.post(url, headers=headers, json=payload, timeout=timeout)
    if r.status_code != 200:
        app.logger.warning("Gemini %s: %s", r.status_code, r.text[:300])
        return None
    candidates = r.json().get("candidates", [])
    if candidates:
        parts = candidates[0].get("content", {}).get("parts", [])
        if parts:
            return parts[0].get("text", "").strip() or None
    return None


def try_model(text, model):
    """Riprova ogni 3 secondi per max 15 secondi con un solo modello."""
    start = time.time()
    while True:
        try:
            result = call_gemini_once(text, model)
        except Exception:
            app.logger.exception("Errore chiamata Gemini (%s)", model)
            result = None
        if result:
            return result
        if time.time() - start + GEMINI_RETRY_EVERY >= GEMINI_MAX_WAIT:
            return None
        time.sleep(GEMINI_RETRY_EVERY)


def process_with_gemini(text):
    """Modello principale per 15 secondi, poi modello di riserva per altri 15.
    Restituisce None se nessuno dei due riesce."""
    if not GEMINI_API_KEY:
        return None
    models = [GEMINI_MODEL]
    if GEMINI_FALLBACK_MODEL and GEMINI_FALLBACK_MODEL != GEMINI_MODEL:
        models.append(GEMINI_FALLBACK_MODEL)
    for model in models:
        result = try_model(text, model)
        if result:
            return result
        app.logger.warning("Modello %s non ha risposto, passo al successivo", model)
    return None


# ---------------------------------------------------------------- TELEGRAM

def tg(method, **payload):
    try:
        requests.post(f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/{method}", json=payload, timeout=10)
    except Exception:
        app.logger.exception("Errore Telegram %s", method)


def send(chat_id, text, buttons=None):
    payload = {"chat_id": chat_id, "text": text}
    if buttons:
        payload["reply_markup"] = {"inline_keyboard": [[{"text": t, "callback_data": d} for t, d in buttons]]}
    tg("sendMessage", **payload)


def handle_message(msg):
    chat_id = msg["chat"]["id"]
    text = (msg.get("text") or "").strip()
    if not text:
        send(chat_id, "Invia la nota come messaggio di testo.")
        return
    if text.startswith("/start"):
        send(chat_id, "Ciao! Scrivimi la nota: la sistemo e, se ti va bene, la pubblichi sul diario.")
        return

    user = msg.get("from", {})
    name = user.get("first_name") or user.get("username") or "Operatore"
    register_operator(chat_id, name)

    pending = PENDING_EDITS.get(chat_id)
    if pending and pending.get("awaiting_edit"):
        PENDING_EDITS[chat_id] = {"text": text, "operator": name, "awaiting_edit": False}
        send(chat_id, f"Ecco la nota corretta:\n\n{text}",
             [("Pubblica", "publish"), ("Modifica", "edit")])
        return

    send(chat_id, "Elaboro la nota...")
    cleaned = process_with_gemini(text)
    if cleaned:
        PENDING_EDITS[chat_id] = {"text": cleaned, "operator": name, "awaiting_edit": False}
        send(chat_id, f"Ecco la nota pronta:\n\n{cleaned}",
             [("Pubblica", "publish"), ("Modifica", "edit")])
    else:
        PENDING_EDITS[chat_id] = {"text": text, "operator": name, "awaiting_edit": False}
        send(chat_id, f"Non sono riuscito a elaborare la nota. Questo è il testo originale:\n\n{text}\n\n"
                      "Premi Modifica per correggerlo a mano.",
             [("Modifica", "edit"), ("Annulla", "cancel")])


def handle_callback(cq):
    chat_id = cq["message"]["chat"]["id"]
    data = cq.get("data")
    tg("answerCallbackQuery", callback_query_id=cq["id"])

    pending = PENDING_EDITS.get(chat_id)
    if not pending:
        send(chat_id, "Non c'è nessuna nota in attesa. Scrivimi una nuova nota.")
        return

    if data == "publish":
        save_entry(now_str(), pending["operator"], pending["text"])
        PENDING_EDITS.pop(chat_id, None)
        send(chat_id, "Nota pubblicata sulla pergamena ✅")
    elif data == "edit":
        pending["awaiting_edit"] = True
        send(chat_id, f"Testo attuale:\n\n{pending['text']}\n\nScrivimi il testo corretto.")
    elif data == "cancel":
        PENDING_EDITS.pop(chat_id, None)
        send(chat_id, "Nota annullata.")


@app.route("/webhook", methods=["POST"])
def webhook():
    try:
        update = request.get_json(silent=True) or {}
        if "callback_query" in update:
            handle_callback(update["callback_query"])
        elif "message" in update:
            handle_message(update["message"])
    except Exception:
        app.logger.exception("Errore nel webhook")
    return "OK", 200


# ---------------------------------------------------------------- PERGAMENA

PERGAMENA_HTML = """<!DOCTYPE html>
<html lang="it">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Diario Operativo</title>
    <style>
        body { background-color: #f4ecd8; font-family: 'Georgia', serif; color: #2c221e; margin: 0; padding: 15px; font-size: 16px; }
        .container { max-width: 850px; margin: 0 auto; background: #fff8eb; border: 2px solid #d4c3a3; box-shadow: 0 4px 15px rgba(0,0,0,0.1); padding: 20px; border-radius: 8px; }
        h1 { text-align: center; color: #5c4033; border-bottom: 2px solid #d4c3a3; padding-bottom: 15px; margin-bottom: 25px; font-size: 1.8em; }
        .filters { display: flex; justify-content: center; gap: 10px; flex-wrap: wrap; margin-bottom: 25px; }
        .filters input, .filters select { padding: 10px; border: 1px solid #d4c3a3; border-radius: 4px; background: #fff; font-family: 'Georgia', serif; font-size: 1em; }
        .filters input[type=text] { width: 40%; min-width: 160px; }
        .entry { margin-bottom: 20px; padding: 15px; border-radius: 6px; border: 1px solid rgba(0,0,0,0.08); }
        .entry-meta { font-size: 0.95em; color: #555; margin-bottom: 10px; font-weight: bold; display: flex; justify-content: space-between; flex-wrap: wrap; gap: 5px; }
        .entry-content { font-size: 1.15em; line-height: 1.6; white-space: pre-wrap; margin-bottom: 15px; }
        .entry-actions { display: flex; gap: 10px; }
        .entry-actions form { margin: 0; }
        .btn-small { background-color: #8b5a2b; color: white; border: none; padding: 8px 14px; font-size: 0.95em; border-radius: 4px; cursor: pointer; text-decoration: none; }
        .btn-small:hover { background-color: #5c4033; }
        .btn-delete { background-color: #a94442; }
        .btn-delete:hover { background-color: #761c19; }
        .actions { text-align: center; margin-top: 30px; display: flex; justify-content: center; gap: 15px; flex-wrap: wrap; }
        .btn { background-color: #8b5a2b; color: white; padding: 12px 20px; text-decoration: none; border-radius: 4px; font-family: sans-serif; font-weight: bold; font-size: 1em; border: none; cursor: pointer; }
        .btn:hover { background-color: #5c4033; }
        .empty { text-align: center; color: #777; font-style: italic; }
    </style>
    <script>
        function copyText(id) {
            const text = document.getElementById('content-' + id).innerText;
            navigator.clipboard.writeText(text).then(() => { alert('Copiato!'); });
        }
    </script>
</head>
<body>
    <div class="container">
        <h1>Diario Operativo</h1>
        <form method="GET" action="/" class="filters">
            <input type="text" name="q" value="{{ search_query }}" placeholder="Cerca nel testo...">
            <select name="operator">
                <option value="">Tutti gli operatori</option>
                {% for name in operators %}
                <option value="{{ name }}" {% if name == selected_operator %}selected{% endif %}>{{ name }}</option>
                {% endfor %}
            </select>
            <button type="submit" class="btn">Filtra</button>
        </form>

        {% if entries %}
            {% for e in entries %}
            <div class="entry" style="background: {{ e.color }};">
                <div class="entry-meta">
                    <span>{{ e.operator }}</span>
                    <span>{{ e.timestamp }}</span>
                </div>
                <div class="entry-content" id="content-{{ e.id }}">{{ e.text }}</div>
                <div class="entry-actions">
                    <button type="button" class="btn-small" onclick="copyText({{ e.id }})">Copia</button>
                    <form method="POST" action="/delete/{{ e.id }}" onsubmit="return confirm('Eliminare questa nota?');">
                        <input type="hidden" name="q" value="{{ search_query }}">
                        <input type="hidden" name="operator" value="{{ selected_operator }}">
                        <button type="submit" class="btn-small btn-delete">Elimina</button>
                    </form>
                </div>
            </div>
            {% endfor %}
        {% else %}
            <p class="empty">Nessuna nota trovata.</p>
        {% endif %}

        <div class="actions">
            <a class="btn" href="{{ url_for('export_docx', q=search_query, operator=selected_operator) }}">Esporta in Word</a>
        </div>
    </div>
</body>
</html>
"""

PALETTE = ["#fdf3dc", "#eef4e4", "#e9f0f7", "#f7e9ee", "#f1ecf8", "#fbeee3"]


def color_for(name):
    return PALETTE[sum(ord(c) for c in name) % len(PALETTE)]


@app.route("/")
def index():
    search_query = request.args.get("q", "").strip()
    selected_operator = request.args.get("operator", "").strip()
    entries = load_entries(search_query or None, selected_operator or None)
    for e in entries:
        e["color"] = color_for(e["operator"])
    return render_template_string(
        PERGAMENA_HTML,
        entries=entries,
        operators=load_operator_names(),
        search_query=search_query,
        selected_operator=selected_operator,
    )


@app.route("/delete/<int:entry_id>", methods=["POST"])
def delete(entry_id):
    delete_entry(entry_id)
    return redirect(url_for("index", q=request.form.get("q", ""), operator=request.form.get("operator", "")))


@app.route("/export")
def export_docx():
    from docx import Document
    from docx.shared import Pt

    search_query = request.args.get("q", "").strip()
    selected_operator = request.args.get("operator", "").strip()
    entries = load_entries(search_query or None, selected_operator or None)
    entries.reverse()  # dal più vecchio al più recente

    doc = Document()
    doc.add_heading("Diario Operativo", level=0)
    for e in entries:
        meta = doc.add_paragraph()
        run = meta.add_run(f"{e['timestamp']} - {e['operator']}")
        run.bold = True
        run.font.size = Pt(10)
        doc.add_paragraph(e["text"])

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    filename = f"diario_operativo_{datetime.now(ITALY_TZ).strftime('%Y%m%d')}.docx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
